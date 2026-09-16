from pathlib import Path
import hashlib,json,math,warnings,zipfile
import openpyxl
warnings.simplefilter('ignore',UserWarning)
HERE=Path(__file__).parent
src=HERE/'man-group-run/source.xlsx'
dst=HERE/'operations-run/man-group-operations.xlsx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
a=openpyxl.load_workbook(src,data_only=False)
b=openpyxl.load_workbook(dst,data_only=False)
av=openpyxl.load_workbook(src,data_only=True)
bv=openpyxl.load_workbook(dst,data_only=True)
def same(x,y):
    if isinstance(x,(int,float)) and isinstance(y,(int,float)):return math.isclose(x,y,rel_tol=1e-10,abs_tol=1e-8)
    return x==y
changes=[];cache=[];formats=[]
for s in a:
    for row in s:
        for c in row:
            d=b[s.title][c.coordinate]
            if not same(c.value,d.value):changes.append([s.title,c.coordinate,c.value,d.value])
            if not same(av[s.title][c.coordinate].value,bv[s.title][c.coordinate].value):cache.append([s.title,c.coordinate])
            if c.number_format!=d.number_format:formats.append([s.title,c.coordinate,c.number_format,d.number_format])
with zipfile.ZipFile(dst) as z:
    names=z.namelist()
    native={'zipIntegrityError':z.testzip(),'formalTableFiles':[n for n in names if n.startswith('xl/tables/') and n.endswith('.xml')],'chartFiles':[n for n in names if '/charts/chart' in n and n.endswith('.xml')],'annotationPartFiles':[n for n in names if 'comment' in n.lower() and n.endswith('.xml')]}
facts={'sourceSha256':sha(src),'downloadOriginalUnchanged':sha(src)==sha(Path('/Users/ruggerogargiulo/Downloads/man-group-fees-excel.xlsx')),'originalSheetOrderPreserved':b.sheetnames[:len(a.sheetnames)]==a.sheetnames,'newSheets':b.sheetnames[len(a.sheetnames):],'originalFormulaOrLiteralChanges':changes,'originalCachedValueChanges':cache,'originalNumberFormatChanges':formats,'nativePackage':native,'feeLab':{'margin':bv['Fee lab']['C4'].value,'quarterOne':bv['Fee lab']['F9'].value,'annualFees':bv['Fee lab']['C15'].value},'engine':'Independent file read using openpyxl 3.1.5 and ZIP/XML; Microsoft Excel application not opened.'}
(HERE/'operations-run/independent-verification.json').write_text(json.dumps(facts,indent=2))
print(json.dumps(facts,indent=2))
